"""Build an editable PowerPoint counterpart of the finished LaTeX presentation.

Requires python-pptx and Pillow. Run from any working directory:
    python Presentation/build_powerpoint.py

Text, tables, layout elements and the facility chart are native PowerPoint
objects. The original model plots and literature diagram are image assets.
"""

from pathlib import Path
import re

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
FIGURES = HERE / "figures"
OUTPUT = HERE / "presentation.pptx"

WIDTH, HEIGHT = 13 + 1 / 3, 7.5
GREEN = "00843D"
INK = "17352D"
MUTED = "60716B"
MIST = "F0F5F2"
RULE = "D7E3DC"
WHITE = "FFFFFF"
FONT = "Arial"


def logistic_description():
    """Keep this slide's explanatory wording sourced directly from LaTeX."""
    latex = (HERE / "presentation.tex").read_text(encoding="utf-8")
    match = re.search(r"\\textbf\{Logistic:\}\s*(.*?)\\par", latex, re.DOTALL)
    if match is None:
        raise ValueError("Could not find the Logistic description in presentation.tex")
    return " ".join(match.group(1).replace(r"\&", "&").split())


def colour(value):
    return RGBColor.from_string(value)


def flat(shape):
    """Override the default Office theme's shadows for a flat Beamer-style look."""
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def rectangle(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour(fill)
    shape.line.fill.background()
    return flat(shape)


def line(slide, x1, y1, x2, y2, fill=RULE, width=1):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = colour(fill)
    shape.line.width = Pt(width)
    return flat(shape)


def text(slide, x, y, w, h, value, size=17, bold=False, fill=INK,
         align=PP_ALIGN.LEFT, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or value.split("\n")[0][:90]
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, content in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = content
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.10
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = colour(fill)
    return shape


def label(slide, x, y, w, value):
    return text(slide, x, y, w, 0.28, value.upper(), 12, True, GREEN)


def source(slide, x, y, w, value, h=0.42):
    return text(slide, x, y, w, h, value, 10.5, fill=MUTED)


def point(slide, x, y, w, title, body, h=1.13, size=17):
    shape = text(slide, x, y, w, h, "", size)
    shape.name = title
    frame = shape.text_frame
    heading = frame.paragraphs[0]
    heading.text = title
    heading.font.bold = True
    heading.font.size = Pt(size + 0.5)
    heading.space_after = Pt(4)
    paragraph = frame.add_paragraph()
    paragraph.text = body
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = colour(INK)
    paragraph.line_spacing = 1.10
    paragraph.space_after = Pt(0)
    return shape


def takeaway(slide, value, y=6.42, h=0.58, x=0.60, w=12.13, size=16):
    rectangle(slide, x, y, w, h, MIST)
    return text(slide, x + 0.15, y + 0.105, w - 0.30, h - 0.16,
                value, size, True)


def metric(slide, x, y, w, value, title, detail, number_size=40):
    text(slide, x, y, w, 0.75, value, number_size, True, GREEN)
    text(slide, x, y + 0.79, w, 0.31, title, 17, True)
    text(slide, x, y + 1.15, w, 0.42, detail, 11.5, fill=MUTED)


def image(slide, name, x, y, w, h):
    path = FIGURES / name
    with Image.open(path) as asset:
        ratio = asset.width / asset.height
    actual_w = min(w, h * ratio)
    actual_h = actual_w / ratio
    shape = slide.shapes.add_picture(
        str(path), Inches(x + (w - actual_w) / 2), Inches(y),
        width=Inches(actual_w), height=Inches(actual_h)
    )
    shape.name = name
    shape._element.nvPicPr.cNvPr.set("descr", name.replace("_", " ").removesuffix(".png"))
    return shape


def table(slide, x, y, widths, rows, row_heights=None, font_size=15):
    """Native, editable table with restrained booktabs-style horizontal rules."""
    if row_heights is None:
        row_heights = [0.40] * len(rows)
    total_w, total_h = sum(widths), sum(row_heights)
    shape = slide.shapes.add_table(
        len(rows), len(widths), Inches(x), Inches(y), Inches(total_w), Inches(total_h)
    )
    shape.name = "Editable data table"
    native = shape.table
    for index, width in enumerate(widths):
        native.columns[index].width = Inches(width)
    for row_index, row in enumerate(rows):
        native.rows[row_index].height = Inches(row_heights[row_index])
        for col_index, value in enumerate(row):
            cell = native.cell(row_index, col_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = colour(WHITE)
            cell.margin_left = Inches(0.02)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = FONT
                paragraph.font.size = Pt(font_size)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = colour(INK)
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.08
                paragraph.alignment = PP_ALIGN.LEFT if col_index == 0 else PP_ALIGN.RIGHT
            properties = cell._tc.get_or_add_tcPr()
            for edge in ("lnL", "lnR", "lnT", "lnB"):
                border = OxmlElement(f"a:{edge}")
                border.append(OxmlElement("a:noFill"))
                properties.append(border)
    line(slide, x, y, x + total_w, y, INK, 1.1)
    line(slide, x, y + row_heights[0], x + total_w, y + row_heights[0], INK, 0.7)
    line(slide, x, y + total_h, x + total_w, y + total_h, INK, 1.1)
    return native


def base(prs, title, section, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colour(WHITE)
    text(slide, 0.60, 0.24, 12.12, 0.44, title, 25, True)
    line(slide, 0, 0.82, WIDTH, 0.82, RULE, 0.8)
    if number is not None:
        line(slide, 0, 0.82, WIDTH * number / 13, 0.82, GREEN, 1.3)
        text(slide, 12.08, 7.18, 0.65, 0.20, f"{number}/13", 10, align=PP_ALIGN.RIGHT)
    text(slide, 0.60, 7.18, 10.9, 0.20,
         f"Machine Learning in Finance  /  {section}", 10, fill=MUTED)
    return slide


def facility_chart(slide):
    data = CategoryChartData()
    data.categories = ["SBA", "EFF", "ECF"]
    data.add_series("Interrupted", [0.64, 0.56, 0.38])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.68), Inches(1.65),
        Inches(5.18), Inches(2.62), data
    ).chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(14)
    plot = chart.plots[0]
    plot.gap_width = 90
    plot.has_data_labels = True
    plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = "0%"
    plot.data_labels.font.name = FONT
    plot.data_labels.font.size = Pt(15)
    plot.data_labels.font.color.rgb = colour(INK)
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = colour(GREEN)
    plot.series[0].format.line.fill.background()
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 0.8
    chart.value_axis.major_unit = 0.2
    chart.value_axis.tick_labels.number_format = "0%"
    chart.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = colour(RULE)
    for axis in (chart.value_axis, chart.category_axis):
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(13)
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.fill.background()
    scaling = chart.category_axis._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling")
    orientation = scaling.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}orientation")
    if orientation is None:
        orientation = OxmlElement("c:orientation")
        scaling.append(orientation)
    orientation.set("val", "maxMin")
    return chart


def build():
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    prs.core_properties.title = "Reading IMF Programs on Day One"
    prs.core_properties.subject = "Approval-time prediction of IMF programme interruptions"
    prs.core_properties.author = (
        "Mathis Clément; Ethan Moore; Ali Haidar; Tristan Zemp; "
        "Joao Matteo Pisaturo; José Pablo Caldas"
    )
    prs.core_properties.keywords = "IMF; machine learning; University of St. Gallen"
    prs.core_properties.comments = "Editable PowerPoint counterpart of Presentation/presentation.tex."

    # 1. Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(s, 10.22, 0, WIDTH - 10.22, HEIGHT, MIST)
    rectangle(s, 0, 0, 0.13, HEIGHT, GREEN)
    image(s, "hsg_logo.png", 0.77, 0.65, 3.20, 0.72)
    text(s, 0.77, 2.02, 9.13, 1.45, "Reading IMF Programs\non Day One", 40, True)
    text(s, 0.77, 3.83, 9.05, 0.90,
         "Machine learning on 173 IMF lending programmes\n2002–2016", 20)
    for x, first, second in (
        (0.77, "Mathis Clément", "Tristan Zemp"),
        (3.15, "Ethan Moore", "Joao Matteo Pisaturo"),
        (6.24, "Ali Haidar", "José Pablo Caldas"),
    ):
        text(s, x, 5.66, 3.0, 0.69, f"{first}\n{second}", 14.5)
    text(s, 0.77, 6.67, 9.0, 0.30, "University of St. Gallen  •  7 October 2026", 14, fill=MUTED)
    metric(s, 10.70, 1.52, 2.22, "173", "Programmes", "2002–2016", 43)
    metric(s, 10.70, 3.44, 2.22, "13", "Predictors", "Known at approval", 43)
    metric(s, 10.70, 5.36, 2.22, "3", "Models", "One prediction task", 43)

    # 2. Executive summary
    s = base(prs, "Executive summary", "Overview", 1)
    label(s, 0.60, 1.15, 7.25, "The question")
    text(s, 0.60, 1.53, 7.32, 0.82,
         "Can information available at approval predict permanent interruption?", 22)
    point(s, 0.60, 2.60, 7.28, "A modest signal in development",
          "The random forest reaches about 0.63 balanced accuracy out of fold.")
    point(s, 0.60, 3.79, 7.28, "No reliable test-set advantage",
          "With only 28 later programmes, performance is unstable and close to chance.")
    point(s, 0.60, 4.98, 7.28, "Programme design is most informative",
          "Scheduled reviews, facility type and loan access lead the forest’s importance ranking.")
    label(s, 8.61, 1.15, 4.12, "Road map")
    agenda = ["Literature and research question", "Study goal", "Data and feature engineering",
              "Three prediction models", "Model comparison", "Limitations and conclusion", "Sources"]
    for i, entry in enumerate(agenda):
        text(s, 8.61, 1.62 + i * 0.57, 4.12, 0.48, f"{i + 1}.  {entry}", 16)
    takeaway(s, "Approval-day information contains some signal, but reliable out-of-sample prediction remains difficult.")

    # 3. Literature and question
    s = base(prs, "Why predict interruption at approval?", "Research question", 2)
    point(s, 0.60, 1.19, 6.18, "Financial relevance",
          "Interruptions can weaken investor confidence and sovereign financing credibility.")
    point(s, 0.60, 2.50, 6.18, "Politics and conditionality",
          "Political incentives and institutional capacity affect implementation; evidence on structural conditions is mixed.", 1.38)
    source(s, 0.60, 3.91, 6.18, "Ivanova et al. (2003); Joyce (2006); Reinsberg et al. (2022).")
    point(s, 0.60, 4.46, 6.18, "The gap in machine learning",
          "Existing studies predict IMF participation. We ask whether approval-time information predicts permanent interruption.", 1.40)
    source(s, 0.60, 5.99, 6.18, "Agbloyor et al. (2023); Batsuuri et al. (2024).", 0.20)
    label(s, 7.25, 1.19, 5.48, "Timing of IMF programme disbursements")
    image(s, "disbursement_timeline.png", 7.25, 1.92, 5.48, 2.38)
    source(s, 7.25, 4.36, 5.48, "Source: Aiyar and Patnam (2021), IMF Working Paper 21/146, Figure 2, p. 8.")
    text(s, 7.25, 5.07, 5.48, 1.05,
         "A programme can end before its scheduled final review. The prediction must be made at Board approval.", 17)
    takeaway(s, "Can information available at approval predict permanent interruption?")

    # 4. Study goal and sample
    s = base(prs, "Study goal and sample", "Study design", 3)
    label(s, 0.60, 1.18, 5.90, "Binary classification")
    text(s, 0.60, 1.62, 5.90, 0.34, "Outcome: programme interruption", 18, True)
    point(s, 0.85, 2.24, 5.65, "Y = 1",
          "At least one scheduled review was never completed.", 1.07, 18)
    point(s, 0.85, 3.42, 5.65, "Y = 0", "Otherwise.", 0.78, 18)
    text(s, 0.60, 4.36, 5.90, 0.47, "Pr(Y = 1 | X) = f(X),   X = (X₁, …, X₁₃).", 18)
    text(s, 0.60, 5.00, 5.90, 0.88,
         "Deal design, conditions, track record and macroeconomic information observed at approval.", 17)
    takeaway(s, "Purpose: prediction, not causal explanation.", 6.13, 0.55, 0.60, 5.90, 15)
    source(s, 0.60, 6.80, 5.90, "Predictive framing: Shmueli (2010); James et al. (2013).", 0.20)
    label(s, 7.05, 1.18, 5.68, "Which programmes enter the sample?")
    selection = [
        ("1. Review-based facilities", "SBA, EFF, ECF and SCF, including predecessors."),
        ("2. Approval in 2002–2016", "MONA review records start in 2002."),
        ("3. Money at stake", "Non-precautionary arrangements with positive drawings."),
        ("4. Observed outcome", "Matched review records."),
        ("5. Blends merged", "Same country and approval date count as one programme."),
    ]
    for i, (heading, body) in enumerate(selection):
        point(s, 7.05, 1.70 + i * 0.91, 5.68, heading, body, 0.96, 16)
    source(s, 7.05, 6.47, 5.68,
           "Earlier risk flags could help calibrate conditionality and inform sovereign-risk pricing.")

    # 5. Data construction
    s = base(prs, "From raw records to programme-level data", "Data and features", 4)
    label(s, 0.60, 1.18, 7.32, "Data sources and their roles")
    rows = [
        ["Source", "Information collected"],
        ["IMF commitments", "Dates, approved and drawn amounts, access as a share of quota, facility type."],
        ["MONA reviews", "Scheduled and completed reviews; construction of the outcome."],
        ["MONA conditions and purchases", "Structural conditions, prior actions and candidate front-loading measures."],
        ["MONA economic indicators", "Approval-snapshot growth, inflation and current account."],
        ["WEO / WDI", "WEO (Apr. 2026): debt and net lending. WDI: growth / inflation fallback considered. Final macro inputs use MONA."],
    ]
    native = table(s, 0.60, 1.68, [2.17, 5.13], rows, [0.44, 0.87, 0.74, 0.90, 0.85, 1.15], 15)
    for row in native.rows:
        for cell in row.cells:
            cell.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    for row_index in range(1, len(rows)):
        native.cell(row_index, 0).text_frame.paragraphs[0].font.bold = True
    source(s, 0.60, 6.73, 7.31,
           "IMF (2026a–c); World Bank (2026). Revised WEO fiscal data are excluded from the prediction inputs.")
    label(s, 8.55, 1.18, 4.18, "Construction pipeline")
    metric(s, 8.55, 1.81, 4.18, "1,576", "Raw arrangements", "Five raw files; no shared identifier")
    point(s, 8.55, 3.53, 4.18, "Link",
          "By country and closest approval date within ±60 days.", 1.08, 17)
    point(s, 8.55, 4.75, 4.18, "Merge blends",
          "Compare approval-day review schedules with completed reviews.", 1.13, 17)
    text(s, 8.55, 6.00, 4.18, 0.64, "173", 36, True, GREEN)
    text(s, 9.68, 6.10, 3.05, 0.32, "Programme observations", 15, True)
    source(s, 8.55, 6.68, 4.18, "82 interrupted / 91 completed", 0.22)

    # 6. Features
    s = base(prs, "Thirteen inputs, all known at approval", "Data and features", 5)
    groups = [
        (0.60, 1.21, "01 / Deal design · 5 inputs",
         "Log access (% of quota); planned duration; scheduled reviews; concessional facility; extended facility."),
        (0.60, 3.03, "02 / Conditionality · 2 inputs",
         "Structural conditions and prior actions at approval."),
        (7.04, 1.21, "03 / Track record · 3 inputs",
         "Prior programmes over 20 years; time since the previous programme; no previous track record."),
        (7.04, 3.03, "04 / Macroeconomy · 3 inputs",
         "Real GDP growth; CPI inflation; current account as a share of GDP, from the approval snapshot."),
    ]
    for x, y, heading, body in groups:
        label(s, x, y, 5.69, heading)
        text(s, x, y + 0.44, 5.69, 1.11, body, 18)
    rectangle(s, 0.60, 4.96, 12.13, 1.66, MIST)
    label(s, 0.78, 5.15, 11.77, "Twelve candidate entries excluded")
    exclusions = [
        ("6 look-ahead", "Post-approval outcomes and revised data"),
        ("4 redundant / fragile", "Duplicate design or history measures"),
        ("1 missing-data", "Fiscal indicators with too many gaps"),
        ("1 identifier group", "Country and approval dates"),
    ]
    for i, (heading, body) in enumerate(exclusions):
        point(s, 0.78 + i * 2.97, 5.59, 2.76, heading, body, 0.90, 15.5)
    source(s, 0.60, 6.79, 12.13,
           "Output: interrupted (0/1). Feature definitions and exclusion audit: Data_Prep/01_data.R.", 0.20)

    # 7. Exploration
    s = base(prs, "Exploration: programme design stands out", "Data and features", 6)
    label(s, 0.60, 1.20, 5.25, "Interruption rate by facility")
    facility_chart(s)
    source(s, 0.60, 4.53, 5.26,
           "Selected facilities; rounded rates reported in the original presentation.")
    text(s, 0.60, 5.25, 5.26, 0.89,
         "Interruption rises from 38% in 2002–07 to 61% in 2014–16.", 18, True)
    point(s, 6.41, 1.20, 6.32, "More reviews, larger programmes",
          "Interrupted programmes have 7.3 scheduled reviews versus 5.9; access is larger and concessional / extended facilities are less common (p ≤ 0.004).", 1.55)
    point(s, 6.41, 2.99, 6.32, "Weak individual correlations",
          "Reviews: 0.31; access: 0.25; concessional: −0.22. Conditions, macro and track record show no clear outcome differences.", 1.42)
    point(s, 6.41, 4.67, 6.32, "Overlap and non-linearity",
          "Concessional–extended correlation: 0.72. Interruption across access quartiles: 31 / 52 / 41 / 67%.", 1.22)
    text(s, 6.41, 6.02, 6.32, 0.67,
         "These patterns motivate testing flexible models alongside a linear log-odds benchmark.", 16.5)
    source(s, 0.60, 6.82, 12.13,
           "Descriptive evidence: full sample, n = 173; these associations are not causal effects.", 0.20)

    # 8. Forest tuning
    s = base(prs, "Random forest: tuning is relatively flat", "Random forest", 7)
    text(s, 0.60, 1.15, 12.13, 0.38,
         "24 settings: 6 mtry values × 4 tree counts; balanced accuracy under three validation schemes.", 17)
    image(s, "rf_tuning.png", 0.72, 1.73, 11.89, 4.63)
    text(s, 0.60, 6.40, 5.90, 0.65,
         "Flat grid. Curves cross; scores stay within roughly five points. No setting dominates.", 15.5)
    text(s, 7.00, 6.40, 5.73, 0.65,
         "Selected with country CV: 250 trees; mtry = 3; balanced accuracy 0.620.", 15.5)
    s.notes_slide.notes_text_frame.text = (
        "Country-grouped validation holds out entire countries, giving the most conservative "
        "tuning scores here. Source: original random-forest tuning figure."
    )

    # 9. Forest performance
    s = base(prs, "Random forest: the test signal is weak", "Random forest", 8)
    label(s, 0.60, 1.18, 6.71, "Development and test performance")
    table(s, 0.60, 1.70, [3.96, 1.22, 1.54], [
        ["", "Accuracy", "Balanced acc."],
        ["Development, out of fold", "0.641", "0.630"],
        ["Test, final forest", "0.500", "0.508"],
        ["Test, always interrupted", "0.607", "0.500"],
    ], [0.46, 0.53, 0.53, 0.53], 15)
    point(s, 0.60, 4.10, 6.71, "Correctly identified on the test set",
          "8 of 17 interrupted programmes; 6 of 11 completed programmes.", 1.13)
    text(s, 0.60, 5.36, 6.71, 0.93,
         "Development scores use out-of-fold predictions; final performance is measured on later programmes.", 17)
    label(s, 8.05, 1.18, 4.68, "Substantial uncertainty")
    metric(s, 8.05, 1.84, 4.68, "0.51–0.64", "Balanced accuracy", "Across 20 random seeds", 36)
    metric(s, 8.05, 3.81, 4.68, "0.38–0.75", "Bootstrap interval",
           "Reported 95% interval; includes 0.50", 36)
    text(s, 8.05, 5.63, 4.68, 0.67,
         "With only 28 test observations, the result is sensitive to random variation.", 16)
    takeaway(s, "The development signal does not translate into a reliable test-set edge over chance.")

    # 10. Importance
    s = base(prs, "Random forest: programme design leads", "Random forest", 9)
    image(s, "rf_importance.png", 0.60, 1.20, 7.25, 5.44)
    source(s, 0.60, 6.78, 7.25,
           "Permutation importance: mean decrease in accuracy after shuffling each predictor.", 0.23)
    label(s, 8.48, 1.20, 4.25, "The strongest predictors")
    for i, entry in enumerate([
        "1.  Number of scheduled reviews", "2.  Extended facility indicator",
        "3.  Loan access (% of quota, log)",
    ]):
        text(s, 8.48, 1.86 + i * 0.91, 4.25, 0.74, entry, 18)
    text(s, 8.48, 4.84, 4.25, 1.16,
         "Macro variables have little importance here; growth and inflation are at or below zero.", 17)
    takeaway(s, "Predictive importance is not a causal effect.", 6.20, 0.79, 8.48, 4.25, 16)

    # 11. Logistic regression
    s = base(prs, "Logistic regression: a transparent benchmark", "Logistic regression", 10)
    label(s, 0.60, 1.18, 5.68, "Test-set performance")
    table(s, 0.60, 1.60, [3.07, 1.30, 1.30], [
        ["", "Logistic", "Forest"], ["Accuracy", "0.536", "0.500"],
        ["Balanced accuracy", "0.570", "0.508"], ["AUC", "0.684", "0.586"],
    ], [0.37, 0.36, 0.36, 0.36], 14)
    image(s, "logit_roc.png", 0.60, 3.29, 5.68, 3.56)
    label(s, 7.04, 1.18, 5.69, "Interpretability")
    image(s, "logit_coefficients.png", 7.04, 1.67, 5.69, 3.88)
    point(s, 7.04, 5.70, 5.69, "Logistic",
          logistic_description(), 0.59, 14.5)
    point(s, 7.04, 6.38, 5.69, "Forest",
          "Captures non-linearities; more flexible, harder to interpret.", 0.59, 14.5)
    s.notes_slide.notes_text_frame.text = (
        "Coefficient magnitudes depend on the units of each predictor; they are not directly "
        "comparable measures of importance. Logistic classification uses a 0.50 threshold."
    )

    # 12. Neural network
    s = base(prs, "Neural network: small is enough", "Neural network", 11)
    label(s, 0.60, 1.18, 7.25, "Country-grouped 10-fold cross-validation")
    image(s, "nn_tuning.png", 0.60, 1.73, 7.20, 4.37)
    source(s, 0.60, 6.20, 7.20,
           "One hidden layer. Size: 1, 2, 3 or 5 neurons; weight decay: 0.01, 0.1, 0.5 or 1.")
    point(s, 8.35, 1.20, 4.38, "A compact network performs best",
          "Two neurons and a light penalty (decay 0.01) give the highest tuning score, about 0.65.", 1.50, 16.5)
    point(s, 8.35, 3.02, 4.38, "More capacity can hurt",
          "With light regularisation, increasing from two to five neurons reduces the score to about 0.54.", 1.46, 16.5)
    point(s, 8.35, 4.87, 4.38, "Heavy regularisation underfits",
          "With decay 1, performance stays near chance (roughly 0.51–0.53).", 1.33, 16.5)
    takeaway(s, "With 145 training programmes, added model complexity does not guarantee better predictions.", 6.62, 0.45, size=15)

    # 13. Model comparison
    s = base(prs, "Model comparison: no reliable winner", "Model comparison", 12)
    image(s, "model_comparison.png", 0.60, 1.30, 7.07, 4.27)
    source(s, 0.60, 5.80, 7.07,
           "Original test-set comparison. The 0.50 line is a chance reference for balanced accuracy and AUC, not the majority-class accuracy benchmark.", 0.51)
    label(s, 8.11, 1.18, 4.62, "All models / same 28 test programmes")
    table(s, 8.11, 1.66, [1.77, 0.95, 0.95, 0.95], [
        ["Model", "Acc.", "BA", "AUC"], ["Logistic", "0.536", "0.570", "0.684"],
        ["Forest", "0.500", "0.508", "0.586"], ["Neural net", "0.500", "0.556", "0.567"],
        ["Always 1", "0.607", "0.500", "0.500"],
    ], [0.40] * 5, 13.5)
    text(s, 8.11, 4.02, 4.62, 0.91,
         "Neither the forest nor the neural network is reliably better than the other.", 17)
    text(s, 8.11, 5.08, 4.62, 1.22,
         "Logistic regression has the highest reported BA and AUC, but this small test sample cannot establish a reliable ranking.", 17)
    takeaway(s, "More flexible models do not demonstrate a dependable advantage in this sample.")

    # 14. Conclusion
    s = base(prs, "Limitations and conclusion", "Conclusion", 13)
    label(s, 0.60, 1.18, 5.72, "What limits the evidence")
    point(s, 0.60, 1.73, 5.72, "Small sample",
          "145 training and 28 test programmes, with only 11 completed. One extra correct prediction changes BA by about 0.03–0.05.", 1.42)
    point(s, 0.60, 3.37, 5.72, "One time split",
          "The interruption share rises from 45% in development to 61% in the test period.", 1.17)
    point(s, 0.60, 4.83, 5.72, "Coarse approval-day inputs",
          "Political commitment, institutional detail and post-approval shocks are not captured.", 1.15)
    label(s, 7.03, 1.18, 5.70, "What we learn")
    point(s, 7.03, 1.73, 5.70, "The signal is weak",
          "About 0.63 out-of-fold BA for the forest, but no reliable edge over chance on the test set.", 1.42)
    point(s, 7.03, 3.37, 5.70, "Deal design predicts best",
          "Reviews, facility type and loan size lead. Macro variables contribute little in this forest.", 1.17)
    point(s, 7.03, 4.83, 5.70, "The metric matters",
          "Plain accuracy can reward an always-interrupted rule. Balanced accuracy gives equal weight to both outcomes.", 1.36)
    takeaway(s, "Approval-time prediction is promising as a research question, but the current evidence is too weak for a dependable risk flag.", 6.34, 0.74, size=16)

    # 15. Literature references
    s = base(prs, "Sources / IMF programmes and implementation", "Sources")
    references = [
        (0.60, 1.22, 1.85, "Agbloyor, E. K., Pan, L., Dwumfour, R. A., & Gyeke-Dako, A. (2023). We are back again! What can artificial intelligence and machine learning models tell us about why countries knock at the door of the IMF? Finance Research Letters, 57, 104244."),
        (0.60, 3.45, 1.30, "Aiyar, S., & Patnam, M. (2021). IMF programs and financial flows to offshore centers. IMF Working Paper No. 21/146. International Monetary Fund."),
        (0.60, 5.12, 1.60, "Batsuuri, T., He, S., Hu, R., Leslie, J., & Lutz, F. (2024). Predicting IMF-supported programs: A machine learning approach. IMF Working Paper No. 24/54. International Monetary Fund."),
        (7.03, 1.22, 1.65, "Ivanova, A., Mayer, W., Mourmouras, A., & Anayiotos, G. (2003). What determines the implementation of IMF-supported programs? IMF Working Paper No. 03/8. International Monetary Fund."),
        (7.03, 3.35, 1.36, "Joyce, J. P. (2006). Promises made, promises broken: A model of IMF program implementation. Economics & Politics, 18(3), 339–365."),
        (7.03, 5.01, 1.75, "Reinsberg, B., Stubbs, T., & Kentikelenis, A. (2022). Compliance, defiance, and the dependency trap: International Monetary Fund program interruptions and their impact on capital markets. Regulation & Governance, 16(4), 1022–1041."),
    ]
    for x, y, h, entry in references:
        text(s, x, y, 5.70, h, entry, 15)

    # 16. Methods and data references
    s = base(prs, "Sources / Methods and data", "Sources")
    label(s, 0.60, 1.18, 5.70, "Methods")
    label(s, 7.03, 1.18, 5.70, "Data sources")
    for y, entry in [
        (1.75, "He, H., & Garcia, E. A. (2009). Learning from imbalanced data. IEEE Transactions on Knowledge and Data Engineering, 21(9), 1263–1284."),
        (3.30, "James, G., Witten, D., Hastie, T., & Tibshirani, R. (2013). An introduction to statistical learning: With applications in R. Springer."),
        (4.85, "Shmueli, G. (2010). To explain or to predict? Statistical Science, 25(3), 289–310."),
    ]:
        text(s, 0.60, y, 5.70, 1.24, entry, 15)
    for y, entry in [
        (1.75, "International Monetary Fund. (2026a). IMF lending commitments [Database]."),
        (2.87, "International Monetary Fund. (2026b). Monitoring of Fund Arrangements (MONA) [Database]."),
        (4.08, "International Monetary Fund. (2026c). World Economic Outlook database: April 2026 edition [Database]."),
        (5.38, "World Bank. (2026). World Development Indicators (WDI) [Database]."),
    ]:
        text(s, 7.03, y, 5.70, 0.96, entry, 15)
    source(s, 0.60, 6.67, 12.13,
           "The WEO and WDI references are retained from the original presentation; the final approval-time macro predictors are from MONA.")

    assert len(prs.slides) == 16
    prs.save(OUTPUT)
    print(f"Created {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
